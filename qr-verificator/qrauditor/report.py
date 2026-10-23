"""Exportación del informe final a CSV y Excel."""
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import STATUS_LABELS, AuditRow

HEADERS = ["Pieza / página", "Página", "URL extraída", "Estado HTTP", "Estado del QR",
           "Tiempo de respuesta (ms)", "URL final", "Redirecciones", "Detalle"]


def _cells(r: AuditRow) -> list:
    return [r.piece, r.page, r.url, r.http_status or "", STATUS_LABELS[r.qr_status],
            r.response_ms if r.response_ms is not None else "", r.final_url, r.redirects, r.detail]


def write_csv(rows: list[AuditRow], path: Path) -> None:
    # utf-8-sig: Excel abre los acentos correctamente.
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(HEADERS)
        w.writerows(_cells(r) for r in rows)


def write_xlsx(rows: list[AuditRow], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Auditoría QR"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4F378B")
    for r in rows:
        ws.append(_cells(r))
        if not r.qr_status.is_success:
            ws.cell(ws.max_row, 5).fill = PatternFill("solid", fgColor="F9DEDC")
    for i, h in enumerate(HEADERS, 1):
        width = max(len(h), *(len(str(c.value or "")) for c in ws[get_column_letter(i)][1:60]), 8)
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 60)
    ws.freeze_panes = "A2"
    wb.save(path)
